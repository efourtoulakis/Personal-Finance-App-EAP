"""
Excel export για τις συναλλαγές συγκεκριμένου μήνα.
Εδώ επιλέγουμε Excel (.xlsx), γιατί είναι εύκολο να ανοιχτεί από τους περισσότερους χρήστες.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from Db.database import select_month_totals, select_transactions
from datetime import datetime
import re


AUTO_REC_MARKER_PATTERN = re.compile(r"\s*\[AUTO-REC:\d+\]\s*")


def date_to_greek(date_text):
    """Μετατρέπει ημερομηνία από YYYY-MM-DD σε DD-MM-YYYY για εμφάνιση στο Excel."""

    try:
        return datetime.strptime(date_text, "%Y-%m-%d").strftime("%d-%m-%Y")
    except (TypeError, ValueError):
        return date_text or ""


def description_to_display(description):
    """
    Κρύβει το εσωτερικό marker recurring από το Excel.

    Το marker μένει στη βάση για αποφυγή διπλοεγγραφών,
    αλλά δεν πρέπει να εμφανίζεται στον τελικό χρήστη.
    """

    if not description:
        return ""

    return AUTO_REC_MARKER_PATTERN.sub("", str(description)).strip()





def export_month_to_excel(month, year, file_path):
    """Δημιουργεί αρχείο Excel με τις συναλλαγές και σύνοψη μήνα."""

    # Διαβάζουμε τα δεδομένα από το database layer.
    transactions = select_transactions(month, year)
    totals = select_month_totals(month, year)



    # Δημιουργούμε νέο workbook.
    wb = Workbook()

    # Παίρνουμε το ενεργό φύλλο.
    ws = wb.active

    # Ονομάζουμε το φύλλο με απλό και κατανοητό όνομα.
    ws.title = "Monthly Report"

    # Τίτλος αναφοράς.
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Αναφορά οικονομικών κινήσεων - {int(month):02d}/{int(year)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Σύνοψη μήνα.
    ws["A3"] = "Σύνολο εσόδων"
    ws["B3"] = totals["income"]
    ws["A4"] = "Σύνολο εξόδων"
    ws["B4"] = -totals["expense"]
    ws["A5"] = "Υπόλοιπο"
    ws["B5"] = totals["balance"]

    # Επικεφαλίδες πίνακα.
    headers = ["ID", "Ημερομηνία", "Κατηγορία", "Τύπος", "Ποσό", "Περιγραφή"]
    start_row = 7

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.font = Font(bold=True)

    # Γράφουμε μία γραμμή για κάθε συναλλαγή.
    for row_index, tran in enumerate(transactions, start=start_row + 1):
        ws.cell(row=row_index, column=1, value=tran["TranID"])
        ws.cell(row=row_index, column=2, value=date_to_greek(tran["TranRegDate"]))
        ws.cell(row=row_index, column=3, value=tran["Category"])
        ws.cell(row=row_index, column=4, value="Έσοδο" if tran["CategoryType"] == "income" else "Έξοδο")
        
        # Το ποσό στη βάση αποθηκεύεται πάντα θετικό.
        # Στο Excel όμως θέλουμε τα έξοδα να εμφανίζονται με αρνητικό πρόσημο.
        amount = float(tran["TranAmount"])

        if tran["CategoryType"] == "expense":
            amount = - amount

        ws.cell(row=row_index, column=5, value=amount)

        ws.cell(row=row_index, column=5, value=amount)
        ws.cell(row=row_index, column=6, value=description_to_display(tran["TranDescr"]))

 

 # Σταθερά πλάτη στηλών για πιο καθαρή εμφάνιση του report.
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 35
    
    
    
    # Αποθηκεύουμε το αρχείο στο path που επέλεξε ο χρήστης.
    wb.save(file_path)

    # Επιστρέφουμε το path ώστε το GUI να εμφανίσει μήνυμα επιτυχίας.
    return file_path
